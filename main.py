"""ISC West exhibitor scraper.

Flow per company (from a text file, one company per line, loop ends when empty):
  1. Put company name in the directory `query` param.
  2. Click the FIRST link that leads to that company (exhibitor-details link).
  3. Scrape the 11 columns matching `ICS West Exhibitors.xlsx`.
"""

from __future__ import annotations

import os
import queue
import sys
import threading
from platform import system
from time import sleep
from urllib.parse import quote_plus

import colorama
from botasaurus.browser import Driver, browser
from colorama import Fore

colorama.init(autoreset=True)

# ---------------------------------------------------------------- config ---

COMPANIES_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "companies.txt")
OUTPUT_XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ICS West Output.xlsx")

DIRECTORY_URL = (
    "https://www.discoverisc.com/west/en-us/2025-exhibitors/"
    "exhibitor-directory.html?locale=en-US&query={query}"
)

HEADERS = (
    "Company Name",
    "Website",
    "email",
    "LinkedIn URL",
    "Why visit our stand",
    "Description",
    "Brands we will feature",
    "Product Categories",
    "Address",
    "Stand(s) #",
    "Source url",
)

# Detail-page selectors (verified live: Napco page renders all of these).
SEL_COMPANY = "#exhibitor-details .details-header h1.wrap-word"
SEL_WEBSITE = '#exhibitor-details a[data-dtm="exhibitorDetails_externalLink"]'
SEL_EMAIL = '#exhibitor-details a[data-dtm="exhibitorDetails_emailLink"]'
SEL_LINKEDIN = '#exhibitor-details .social-media a[href*="linkedin.com"]'
SEL_WHY = "#exhibitor_details_showobjective > p"
SEL_DESC = "#exhibitor_details_description > p"
SEL_BRANDS = "#exhibitor_details_brands > p"
SEL_SHOW_ALL = '#exhibitor-details button[data-dtm="exhibitorDetails_categories_showAllCategories_button"]'
SEL_CAT_TAGS = '#exhibitor-details .categories-section span[data-testid="filter-tag"]'
SEL_OVERFLOW = '#exhibitor-details .categories-section span[data-testid="overflow-count"]'
SEL_ADDRESS = "#exhibitor_details_address p span"
SEL_STAND = "#exhibitor-details .stand-details > span.stand-reference-label"
SEL_DETAIL_H1 = "#exhibitor-details h1.wrap-word"
SEL_RESULT_LINKS = 'a[href*="exhibitor-details."]'

# ------------------------------------------------------------- browser ------

brave_path_in_linux = "/usr/bin/brave"
brave_path_in_windows = (
    "C:\\Program Files\\BraveSoftware\\Brave-Browser\\Application\\brave.exe"
)

browser_kwargs = {
    "chrome_executable_path": (
        brave_path_in_windows if system() == "Windows" else brave_path_in_linux
    ),
    "output": None,
    "close_on_crash": True,
    "raise_exception": True,
    "create_error_logs": False,
}

browser_kwargs["add_arguments"] = [
    "--disable-background-networking",
    "--disable-component-update",
    "--disable-domain-reliability",
    "--disable-sync",
    "--metrics-recording-only",
    "--no-pings",
    "--safebrowsing-disable-auto-update",
    "--disable-renderer-backgrounding",
    "--disable-background-timer-throttling",
    "--disable-backgrounding-occluded-windows",
    "--disable-dev-shm-usage",
]

BLOCKED_URL_PATTERNS = [
    "*.brave.com",
    "*.bravesoftware.com",
    "*.bsg.brave.com",
    "*.pcdn.brave.com",
    "*rewards.brave.com*",
    "*ads.brave.com*",
    "*componentupdater.brave.com*",
    "*go-updater.brave.com*",
    "*.gvt1.com",
    "*.gvt2.com",
    "clients2.google.com",
    "update.googleapis.com",
]

# -------------------------------------------------------------- helpers -----


def _clean(value) -> str:
    return " ".join((value or "").split())


def _text_safe(driver: Driver, selector: str) -> str:
    try:
        el = driver.select(selector, wait=None)
    except Exception:
        return ""
    if not el:
        return ""
    try:
        return _clean(el.text)
    except Exception:
        return ""


def _href_safe(driver: Driver, selector: str) -> str:
    try:
        el = driver.select(selector, wait=None)
    except Exception:
        return ""
    if not el:
        return ""
    try:
        return _clean(el.get_attribute("href") or "")
    except Exception:
        return ""


def _accept_cookies(driver: Driver) -> None:
    try:
        if driver.select("#onetrust-accept-btn-handler", wait=None):
            driver.click("#onetrust-accept-btn-handler")
            sleep(0.5)
    except Exception:
        pass


def _expand_categories(driver: Driver) -> None:
    """Click 'Show all categories' so ALL product categories render."""
    try:
        btn = driver.select(SEL_SHOW_ALL, wait=None)
        if not btn:
            return
        if "show all" not in _clean(btn.text).lower():
            return
        try:
            driver.scroll_into_view(SEL_SHOW_ALL)
        except Exception:
            pass
        driver.click(SEL_SHOW_ALL)
        for _ in range(15):
            sleep(0.2)
            try:
                if not driver.select(SEL_OVERFLOW, wait=None):
                    return
            except Exception:
                return
    except Exception:
        pass


def _categories(driver: Driver) -> str:
    try:
        els = driver.select_all(SEL_CAT_TAGS, wait=None) or []
    except Exception:
        return ""
    names = []
    for el in els:
        try:
            name = _clean(el.text)
        except Exception:
            continue
        if name:
            names.append(name)
    return ", ".join(names)


def _address(driver: Driver) -> str:
    try:
        els = driver.select_all(SEL_ADDRESS, wait=None) or []
    except Exception:
        return ""
    lines = []
    for el in els:
        try:
            line = _clean(el.text)
        except Exception:
            continue
        if line:
            lines.append(line)
    return "\n".join(lines)


def _stand(driver: Driver) -> str:
    try:
        els = driver.select_all(SEL_STAND, wait=None) or []
    except Exception:
        return ""
    numbers = []
    for el in els:
        try:
            value = _clean(el.text).rstrip(":")
        except Exception:
            continue
        if not value or value.lower().startswith("stand"):
            continue
        numbers.append(value)
    return ", ".join(numbers)


def _email(driver: Driver) -> str:
    text = _text_safe(driver, SEL_EMAIL)
    if text:
        return text
    href = _href_safe(driver, SEL_EMAIL)
    if href.lower().startswith("mailto:"):
        return href.split(":", 1)[1].strip()
    return href


def extract_detail(driver: Driver) -> dict:
    """Map the OPEN exhibitor detail DOM to spreadsheet columns."""
    _expand_categories(driver)
    # LinkedIn logo is NOT always the first social logo -> match by href.
    return {
        "Company Name": _text_safe(driver, SEL_COMPANY),
        "Website": _href_safe(driver, SEL_WEBSITE),
        "email": _email(driver),
        "LinkedIn URL": _href_safe(driver, SEL_LINKEDIN),
        "Why visit our stand": _text_safe(driver, SEL_WHY),
        "Description": _text_safe(driver, SEL_DESC),
        "Brands we will feature": _text_safe(driver, SEL_BRANDS),
        "Product Categories": _categories(driver),
        "Address": _address(driver),
        "Stand(s) #": _stand(driver),
        "Source url": (driver.current_url or ""),
    }


def _first_company_link(driver: Driver, timeout_s: float = 30.0):
    """Return the FIRST exhibitor-details link on the search page, else None."""
    waited = 0.0
    while waited < timeout_s:
        try:
            links = driver.select_all(SEL_RESULT_LINKS, wait=None) or []
        except Exception:
            links = []
        for link in links:
            try:
                href = (link.get_attribute("href") or "").strip()
            except Exception:
                continue
            if href and "exhibitor-details." in href:
                return link
        sleep(1.0)
        waited += 1.0
    return None


def scrape_one(driver: Driver, company_name: str) -> dict:
    """Query the directory, open the FIRST company link, extract fields."""
    empty = {h: "" for h in HEADERS}
    empty["Company Name"] = company_name

    search_url = DIRECTORY_URL.format(query=quote_plus(company_name))
    try:
        driver.get(search_url, timeout=60)
    except Exception as exc:
        print(Fore.RED + f"[{company_name}] search load failed: {exc}")
        empty["Source url"] = search_url
        return empty

    sleep(4)
    _accept_cookies(driver)

    link = _first_company_link(driver, timeout_s=30.0)
    if link is None:
        print(Fore.YELLOW + f"[{company_name}] no result link found, skipped.")
        empty["Source url"] = search_url
        return empty

    try:
        href = (link.get_attribute("href") or "").strip()
    except Exception:
        href = ""
    if href and href.startswith("/"):
        href = "https://www.discoverisc.com" + href
    target = href or search_url

    try:
        # Direct GET is more reliable than clicking in this flaky SPA.
        driver.get(target, timeout=90)
    except Exception:
        try:
            link.click()
        except Exception as exc:
            print(Fore.RED + f"[{company_name}] open detail failed: {exc}")
            empty["Source url"] = target
            return empty

    sleep(5)
    _accept_cookies(driver)

    try:
        found = driver.select(SEL_DETAIL_H1, wait=20)
    except Exception:
        found = None
    if not found:
        print(Fore.YELLOW + f"[{company_name}] detail did not render: {target}")
        empty["Source url"] = target
        return empty

    try:
        row = extract_detail(driver)
    except Exception as exc:
        print(Fore.RED + f"[{company_name}] extract failed: {exc}")
        empty["Source url"] = target
        return empty

    if not row.get("Company Name"):
        row["Company Name"] = company_name
    if not row.get("Source url"):
        row["Source url"] = target
    return row


# --------------------------------------------------------------- worker -----

work_queue: queue.Queue = queue.Queue()
results: list = []
results_lock = threading.Lock()


class Scrape:
    def run(self, browser: Driver):
        try:
            browser.block_urls(BLOCKED_URL_PATTERNS)
        except Exception:
            pass
        while True:
            try:
                company_name = work_queue.get_nowait()
            except queue.Empty:
                return
            try:
                print(Fore.CYAN + f"[+] {company_name}")
                row = scrape_one(browser, company_name)
                with results_lock:
                    results.append(row)
                print(Fore.GREEN + f"  done: {row.get('Company Name')}")
            except Exception as exc:
                print(Fore.RED + f"[{company_name}] failed: {exc}")
                with results_lock:
                    results.append(
                        {"Company Name": company_name, "Source url": ""}
                    )
            finally:
                work_queue.task_done()


@browser(**browser_kwargs, wait_for_complete_page_load=True)
def run_with_browser(browser: Driver, selenium_obj: "Scrape"):
    try:
        selenium_obj.run(browser)
    except Exception as error:
        print(Fore.RED + str(error))


# ----------------------------------------------------------------- xlsx ------


def write_xlsx(rows: list, path: str) -> str:
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    if openpyxl is not None:
        from openpyxl.styles import Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(list(HEADERS))
        wrap = Alignment(wrap_text=True)
        addr_idx = HEADERS.index("Address")
        for row in rows:
            if isinstance(row, dict):
                values = [row.get(h, "") or "" for h in HEADERS]
            else:
                values = [row[i] if i < len(row) else "" for i in range(len(HEADERS))]
            ws.append([str(v) for v in values])
            ws.cell(ws.max_row, addr_idx + 1).alignment = wrap
        wb.save(path)
        return path
    # stdlib fallback (no openpyxl installed)
    from xml.sax.saxutils import escape
    from zipfile import ZIP_DEFLATED, ZipFile

    def col_letter(i: int) -> str:
        n = i + 1
        s = ""
        while n:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s

    def xml_text(v) -> str:
        return escape(str(v or "")).replace("\n", "&#10;")

    sheet_rows = [list(HEADERS)]
    for row in rows:
        if isinstance(row, dict):
            sheet_rows.append([row.get(h, "") or "" for h in HEADERS])
        else:
            sheet_rows.append([row[i] if i < len(row) else "" for i in range(len(HEADERS))])
    addr_idx = HEADERS.index("Address")
    parts = []
    for rn, vals in enumerate(sheet_rows, start=1):
        cells = [f'<row r="{rn}">']
        for i, v in enumerate(vals):
            ref = f"{col_letter(i)}{rn}"
            style = ' s="1"' if rn > 1 and i == addr_idx else ""
            cells.append(
                f'<c r="{ref}"{style} t="inlineStr"><is><t xml:space="preserve">{xml_text(v)}</t></is></c>'
            )
        cells.append("</row>")
        parts.append("".join(cells))
    last_col = col_letter(len(HEADERS) - 1)
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f"<dimension ref=\"A1:{last_col}{len(sheet_rows)}\"/>"
        "<sheetData>" + "".join(parts) + "</sheetData></worksheet>"
    )
    styles_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill>'
        '<fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border/></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        "<cellXfs count=\"2\">"
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1">'
        '<alignment wrapText="1"/></xf></cellXfs></styleSheet>'
    )
    wb_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Sheet" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    types_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/></Relationships>'
    )
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
        'Target="styles.xml"/></Relationships>'
    )
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", types_xml)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("xl/workbook.xml", wb_xml)
        z.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        z.writestr("xl/styles.xml", styles_xml)
        z.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return path


def load_companies(path: str) -> list:
    if not os.path.exists(path):
        with open(path, "w", encoding="utf-8") as f:
            f.write("Napco Security Technologies Inc.\n")
        print(Fore.YELLOW + f"Created example {path} - add one company per line.")
    with open(path, encoding="utf-8-sig") as f:
        names = [line.strip() for line in f]
    # Loop ends when the file is empty: drop blanks.
    return [n for n in names if n]


# ------------------------------------------------------------------ main -----

if __name__ == "__main__":
    try:
        num_browsers = int(input("Enter Number Of Threading : ").strip() or "1")
    except Exception:
        num_browsers = 1
    num_browsers = max(1, min(num_browsers, 10))

    companies = load_companies(COMPANIES_FILE)
    if not companies:
        print(Fore.RED + f"No companies found in {COMPANIES_FILE}")
        sys.exit(1)

    print(Fore.LIGHTCYAN_EX + f"[+] {len(companies)} companies queued.")
    for name in companies:
        work_queue.put(name)

    threads = []
    for _ in range(num_browsers):
        t = threading.Thread(target=run_with_browser, args=(Scrape(),))
        threads.append(t)
        t.start()
    for t in threads:
        t.join()

    # Keep output in input order.
    order = {name: i for i, name in enumerate(companies)}
    ordered = sorted(results, key=lambda r: order.get((r.get("Company Name") or ""), 9999))

    write_xlsx(ordered, OUTPUT_XLSX)
    print(" ")
    print(Fore.LIGHTGREEN_EX + f"DONE! {len(ordered)} rows -> {OUTPUT_XLSX}")
    sys.exit('Press Enter To Exit ...')
